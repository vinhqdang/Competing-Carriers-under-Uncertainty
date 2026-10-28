#!/usr/bin/env python3
"""Builds the Word files of paper/submission/ready/ (highlights, cover letter, declaration, title page, CRediT)
from the repository sources. Run after flatten.py has produced ready/reference_upload_with_author/."""
import os, re, subprocess, shutil
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

H = os.path.dirname(os.path.abspath(__file__)); P = os.path.dirname(H); R = f'{H}/ready'
URL = 'https://github.com/vinhqdang/Competing-Carriers-under-Uncertainty'
TITLE = 'City charges for competing delivery carriers: exact best responses under probability-box travel times and leader pricing without equilibrium search'
CREDIT = ('Quang-Vinh Dang: Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, Resources, '
          'Data curation, Writing - original draft, Writing - review & editing, Visualization, Project administration.')
DOI_ = 'The author declares that he has no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.'
FUND = 'This research received no specific grant from any funding agency in the public, commercial, or not-for-profit sectors.'

def newdoc():
    d = Document(); st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12); return d
def para(d, text, bold=False, italic=False, size=None, align=None):
    p = d.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    return p
def md_runs(p, text):
    for i, part in enumerate(re.split(r'\*\*', text)):
        r = p.add_run(part); r.bold = (i % 2 == 1)
def to_pdf(docx):
    subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', os.path.dirname(docx), docx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

# highlights
hl = [l.strip() for l in open(f'{P}/highlights.txt') if l.strip()]
assert 3 <= len(hl) <= 5 and all(len(x) <= 85 for x in hl)
open(f'{R}/highlights.txt', 'w').write('\n'.join(hl) + '\n')
d = newdoc(); para(d, 'Highlights', bold=True, size=14)
for x in hl: d.add_paragraph(x, style='List Bullet')
d.save(f'{R}/highlights.docx')

# cover letter
shutil.copy(f'{H}/cover_letter.md', f'{R}/cover_letter.md')
d = newdoc()
for blk in open(f'{H}/cover_letter.md').read().strip().split('\n\n'):
    for line in blk.split('\n'):
        p = d.add_paragraph(); md_runs(p, line.rstrip())
d.save(f'{R}/cover_letter.docx'); to_pdf(f'{R}/cover_letter.docx')

# declaration of interest
d = newdoc(); para(d, 'Declaration of interest', bold=True, size=14); para(d, DOI_)
para(d, 'Funding', bold=True, size=12); para(d, FUND); d.save(f'{R}/declaration_of_interest.docx')

# CRediT
d = newdoc(); para(d, 'CRediT authorship contribution statement', bold=True, size=14); para(d, CREDIT)
d.save(f'{R}/credit_statement.docx')

# title page
tex = open(f'{R}/reference_upload_with_author/main.tex').read()
ab = re.search(r'\\begin\{abstract\}(.*?)\\end\{abstract\}', tex, re.S).group(1)
kw = re.search(r'\\begin\{keyword\}(.*?)\\end\{keyword\}', tex, re.S).group(1)
def clean(s):
    s = s.replace('\\%', '%').replace('--', '\u2013').replace('{}', '').replace('~', ' ')
    s = re.sub(r'\\sep', ';', s); s = re.sub(r'\s+', ' ', s).replace(' ;', ';').strip()
    assert '\\' not in s and '$' not in s, s
    return s
ab, kw = clean(ab), clean(kw)
txt = subprocess.run(['pdftotext', f'{R}/manuscript_anonymous.pdf', '-'], capture_output=True, text=True).stdout
pages = txt.split('\f'); body = []
for pg in pages:
    if re.search(r'^References\s*$', pg, re.M):
        body.append(pg.split('\nReferences')[0]); break
    body.append(pg)
wc = len(re.findall(r'\S+', ' '.join(body)))
wc = int(round(wc, -2))
d = newdoc()
para(d, TITLE, bold=True, size=15)
para(d, 'Quang-Vinh Dang', bold=True)
para(d, 'British University Vietnam, Hung Yen, Vietnam')
para(d, 'Corresponding author: Quang-Vinh Dang, vinh.dq4@buv.edu.vn')
para(d, 'ORCID: 0000-0002-3877-8024')
para(d, 'Article type: VSI: Omega55: Smart Cities', italic=True)
para(d, 'Abstract', bold=True); para(d, ab)
para(d, 'Keywords', bold=True); para(d, kw)
para(d, 'Word count', bold=True)
para(d, f'About {wc:,} words of main text (title page to the end of the conclusion, including abstract, headings, equations and table text; excluding references and appendices).')
para(d, 'Data and code availability', bold=True); para(d, f'Code, processed data, parameter sources, result files, figures and the table-generating scripts are available at {URL}.')
para(d, 'Declaration of competing interest', bold=True); para(d, DOI_)
para(d, 'Funding', bold=True); para(d, FUND)
para(d, 'CRediT authorship contribution statement', bold=True); para(d, CREDIT)
d.save(f'{R}/title_page.docx'); to_pdf(f'{R}/title_page.docx')
print('word count', wc)

# figures table
os.makedirs(f'{R}/figures_300dpi', exist_ok=True)
TW = 472.03123 / 72.27
figs = [('Figure 1', 'fig01_overview_schematic.png', 1.0, 'TikZ schematic rendered to PNG at 450 dpi'),
        ('Figure 2', 'fig02_arp_calls_vs_gain.png', 0.65, 'experiments/plot_arp.py re-run at 300 dpi'),
        ('Figure 3', 'fig03_imprecision.png', 1.0, 'experiments/plot_results.py (e4) re-run at 300 dpi'),
        ('Figure 4', 'fig04_instruments.png', 1.0, 'experiments/plot_results.py (e2_heat) re-run at 300 dpi'),
        ('Figure 5', 'fig05_equilibria_overview.png', 0.78, 'copied from figures/ (1600 x 1500 px, source 100 dpi; effective dpi at print size is above 300)'),
        ('Figure 6', 'fig06_dip_calls_vs_gain.png', 0.8, 'experiments/plot_dip.py re-run at 300 dpi'),
        ('Figure 7', 'fig07_theory_absorption.png', 1.0, 'experiments/run_theory_checks.py report, figure code at 300 dpi'),
        ('Figure 8', 'fig08_theory_d2.png', 0.9, 'experiments/run_theory_checks.py report, figure code at 300 dpi')]
rows = ['| Figure | File | Size (px) | Printed width (in) | Effective dpi | Origin |', '|---|---|---|---|---|---|']
for n, f, frac, note in figs:
    shutil.copy(f'{R}/anonymous_upload/{f}', f'{R}/figures_300dpi/{f}')
    w, h = Image.open(f'{R}/figures_300dpi/{f}').size
    pw = frac * TW; assert w / pw >= 300, f
    rows.append(f'| {n} | {f} | {w} x {h} | {pw:.2f} | {w/pw:.0f} | {note} |')
open(f'{R}/figures_300dpi/figures_table.md', 'w').write('Printed width = fraction of the text width (6.53 in).\n\n' + '\n'.join(rows) + '\n')
